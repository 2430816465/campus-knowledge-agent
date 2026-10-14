"""Read-only development fact audit. No inference, embedding or teacher approval.

Expected facts are curated from originals, never generated from retrieval output.
Candidate evidence is not scored as a validated answer or student UI success.
"""
import hashlib
import json
import re
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))


def normalized(text):
    return re.sub(r'\s+','',text)


def contains_facts(texts,facts):
    # A fact must occur in one unit, not be invented by concatenating boxes.
    return all(any(normalized(f) in normalized(t) for t in texts) for f in facts)


def failure_stage(*,original_ok,document_exists,searchable,chunks_ok,retrieved_ok,
                  answer_ok,mode,ocr=False):
    if not original_ok:return 'original_or_fixture_unverified'
    if not document_exists:return 'not_ingested'
    if not searchable:return 'partial_failed_or_not_publicly_searchable'
    if not chunks_ok:return 'ocr_omission' if ocr else 'parse_omission'
    if not retrieved_ok:return 'retrieval_omission'
    if not answer_ok:return 'answer_omission'
    return 'candidate_evidence_only' if mode=='evidence_fallback' else 'literal_answer_present_needs_scope_review'


def original_units(case,content):
    suffix=case['format']
    if suffix=='.pdf':
        import pdfplumber
        with pdfplumber.open(BytesIO(content)) as pdf:
            return [(pdf.pages[case['original_page']-1].dedupe_chars(extra_attrs=()).extract_text() or '')]
    if suffix=='.docx':
        with ZipFile(BytesIO(content)) as z:
            tree=ET.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        return [''.join(p.itertext()) for p in tree.findall('.//w:p',ns)]
    if suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(file_contents=content,on_demand=True)
        try:return [str(c.value) for s in book.sheets() for row in range(s.nrows) for c in s.row(row)]
        finally:book.release_resources()
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(content),read_only=True,data_only=True)
        try:return [str(c) for s in book for row in s.iter_rows(values_only=True) for c in row if c is not None]
        finally:book.close()
    if suffix in ('.png','.jpg') and case.get('visually_verified'):
        # Human visual attestation is bound to this exact original SHA below.
        # Do not silently substitute OCR output as independent ground truth.
        return case['facts']
    raise ValueError('No independent original verification for this fixture format')


def main():
    from sqlalchemy import select
    from app.config import DATA
    from app.db import Session,Document,Chunk,timestamp
    from app.campus_qa import campus_answer
    from app.campus_search import retrieve_public
    cases=json.loads((ROOT/'tests/fixtures/basic_fact_sources.json').read_text('utf-8'))
    results=[]
    with Session() as s:
        for case in cases:
            result={'id':case['id'],'format':case['format'],'question':case['question'],
                    'expected_facts':case['facts'],'original_sha256':case['sha256'],
                    'source_url':case['url'],'student_ui_verified':case.get('student_ui_verified',False)}
            path=DATA/'school-archive/blobs'/case['sha256']
            try:
                content=path.read_bytes()
                original_ok=hashlib.sha256(content).hexdigest()==case['sha256'] and contains_facts(original_units(case,content),case['facts'])
                doc=s.scalar(select(Document).where(Document.sha256==case['sha256'],Document.source_url==case['url']))
                texts=[c.text for c in s.scalars(select(Chunk).where(Chunk.document_id==doc.id))] if doc else []
                chunks_ok=contains_facts(texts,case['facts'])
                candidates,_=retrieve_public(s,case['question'],allow_inference=False)
                answer=campus_answer(s,case['question'],timestamp(),use_model=False)
                retrieved_ok=contains_facts([r.get('quote',r.get('text','')) for r in candidates
                                            if r.get('document_id')==(doc.id if doc else None)],case['facts'])
                answer_ok=contains_facts([answer.get('answer','')],case['facts'])
                stage=failure_stage(original_ok=original_ok,document_exists=bool(doc),
                    searchable=bool(doc and doc.visibility=='public' and doc.status in ('pending_review','published')),
                    chunks_ok=chunks_ok,retrieved_ok=retrieved_ok,answer_ok=answer_ok,
                    mode=answer.get('mode'),ocr=case['format'] in ('.png','.jpg'))
                result.update(stage=stage,original_verified=original_ok,document_id=doc.id if doc else None,
                    document_status=doc.status if doc else None,chunk_fact_present=chunks_ok,
                    retrieval_fact_present=retrieved_ok,answer_literal_present=answer_ok,
                    retrieval_candidates=candidates,answer=answer)
            except Exception as exc:
                result.update(stage='audit_error',error_type=type(exc).__name__,error=str(exc)[:200])
                s.rollback()
            results.append(result)
    report={'checked_at':timestamp(),'mode':'CPU_no_inference_development_audit',
            'model_calls':0,'embedding_calls':0,'database_writes':False,'results':results,
            'limits':['Five curated development samples; not independent accuracy benchmark.',
                      'Only PDF first case has actual student-page verification; other CLI audits do not substitute UI.',
                      'Literal presence does not certify dates, relations, current validity or complete answers.']}
    (ROOT/'reports/basic-fact-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps([{'id':r['id'],'stage':r['stage']} for r in results],ensure_ascii=False))


if __name__=='__main__':main()
