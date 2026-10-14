from io import BytesIO
from hashlib import sha256

import pytest
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from sqlalchemy import select

from app.docx_images import extract_docx_images
from app.ingest import parse
from app.model_evidence import prepare_units
from app.db import Chunk, Audit
from tests.test_appeal_retrieval import session_factory, add
from tests.test_convert_archive_doc import module, archived


def raster():
    stream=BytesIO();Image.new('RGB',(40,30),'white').save(stream,format='PNG')
    return stream.getvalue()


def document(text='', copies=1):
    doc=Document()
    if text:doc.add_paragraph(text)
    for _ in range(copies):doc.add_picture(BytesIO(raster()))
    return doc


def payload(doc):
    stream=BytesIO();doc.save(stream);return stream.getvalue()


def ocr(image):
    assert image.shape==(30,40,3)
    return [([[1,2],[38,2],[38,25],[1,25]],'背景白色或蓝色',.98)]


def test_body_reference_citations_and_duplicate_raster_ocr_only_once():
    calls=[]
    blocks,warnings=extract_docx_images(document(copies=2),lambda image:(calls.append(image.shape) or ocr(image)))
    assert len(calls)==1 and len(blocks)==2
    assert blocks[0][1]==blocks[1][1]=='背景白色或蓝色'
    assert 'Word 图像 2' in blocks[1][0] and 'conf=0.980 xy=1,2;38,2;38,25;1,25' in blocks[1][0]
    assert all(len(loc)<=120 for loc,_ in blocks)
    assert any(sha256(raster()).hexdigest() in w for w in warnings)


@pytest.mark.parametrize('kind',['external','crop','rotation','flip'])
def test_unsupported_display_transforms_or_external_relationship_never_call_ocr(kind):
    doc=document();blip=next(doc.element.body.iter(qn('a:blip')))
    if kind=='external':blip.set(qn('r:link'),'https://example.com/private.png')
    elif kind=='crop':
        crop=OxmlElement('a:srcRect');crop.set('l','1000');blip.getparent().append(crop)
    else:
        transform=next(doc.element.body.iter(qn('a:xfrm')))
        transform.set('rot' if kind=='rotation' else 'flipH','1000' if kind=='rotation' else '1')
    def forbidden(_):raise AssertionError('OCR should not run')
    blocks,warnings=extract_docx_images(doc,forbidden)
    assert not blocks and any('待 OCR' in w for w in warnings)
    assert not any('AssertionError' in w for w in warnings)


@pytest.mark.parametrize('box,text,score',[
    ([[0,0]]*4,'低置信度',.49),([[0,0]]*4,'NaN',float('nan')),
    ([[0,0]]*4,'bad',1.1),([[-1,0]]*4,'越界',.9),([[0,0]]*3,'三角',.9),
    ([[0,0]]*4,'x'*1001,.9),([[0,0]]*4,'',.9),
],ids=['low','nan','high','bounds','quad','long','empty'])
def test_invalid_boxes_are_not_searchable(box,text,score):
    blocks,warnings=extract_docx_images(document(),lambda _: [[box,text,score]])
    assert not blocks and any('待 OCR' in w for w in warnings)


def test_pixel_reference_and_box_budgets_preserve_existing_text(monkeypatch):
    import app.docx_images as images
    monkeypatch.setattr(images,'MAX_IMAGES',1)
    blocks,warnings=extract_docx_images(document(copies=2),ocr)
    assert len(blocks)==1 and any('截断' in w for w in warnings)
    monkeypatch.setattr(images,'MAX_PIXELS',1)
    blocks,warnings=extract_docx_images(document(),ocr)
    assert not blocks and any('像素预算' in w for w in warnings)
    monkeypatch.setattr(images,'MAX_PIXELS',25000000)
    monkeypatch.setattr(images,'MAX_BOXES',0)
    blocks,warnings=extract_docx_images(document(),ocr)
    assert not blocks and any('框超出预算' in w for w in warnings)


def test_mixed_word_body_text_survives_ocr_failure_and_header_images_are_not_processed(monkeypatch):
    import app.ingest as ingest
    doc=document('正常正文');doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raster()))
    def failed(_):raise RuntimeError('offline OCR failed')
    monkeypatch.setattr(ingest,'local_ocr',failed)
    _,_,_,blocks,warnings,_=parse(payload(doc),'.docx')
    assert blocks==[('Word 段落 1','正常正文')]
    assert len([w for w in warnings if '待 OCR' in w])==1


def test_image_only_native_repair_keeps_identity_and_rerun_is_readonly(monkeypatch,tmp_path):
    import app.ingest as ingest
    monkeypatch.setattr(ingest,'local_ocr',ocr)
    content=payload(document())
    with session_factory()() as session:
        parent,_=add(session,'公开通知',['附件'])
        doc,_=add(session,'图像Word',[],parent_id=parent.id)
        doc.mime='.docx';doc.status='parse_failed';doc.sha256=sha256(content).hexdigest()
        doc.source_url='https://jwc.njxzc.edu.cn/image.docx'
        session.commit()
        before=(doc.id,doc.sha256,doc.mime,doc.source_url,doc.title,doc.parent_id,doc.version,doc.visibility)
        store=module.BusinessStore(session)
        row=archived(tmp_path,content,'image.docx');row.update(url=doc.source_url,document_id=doc.id)
        def forbidden(*_):raise AssertionError('Word must not run')
        result=module.convert_row(row,store,tmp_path/'school-archive',tmp_path/'out',tmp_path/'helper',converter=forbidden)
        assert result['state']=='native_parsed' and result['word_invoked'] is False
        assert result['original_document_id']==result['converted_document_id']==doc.id
        assert (doc.id,doc.sha256,doc.mime,doc.source_url,doc.title,doc.parent_id,doc.version,doc.visibility)==before
        assert doc.status=='pending_review'
        ids=set(session.scalars(select(Chunk.id).where(Chunk.document_id==doc.id)))
        assert len(ids)==1
        store.parse_native_original(doc,content,'.docx')
        assert ids==set(session.scalars(select(Chunk.id).where(Chunk.document_id==doc.id)))
        audits=list(session.scalars(select(Audit).where(Audit.entity_id==doc.id)))
        assert len(audits)==1 and audits[0].changes['declared_suffix']=='.docx'
        assert audits[0].changes['embedding_requested'] is False


def test_ocr_context_also_marks_prompt_without_inventing_image_geometry():
    row={'id':'one','document_id':'doc','text':'正文','location':'Word 段落 1',
         'context_chunks':[{'id':'two','text':'0.4米','location':'Word 图像 1 / OCR 1'}]}
    _,originals,prompt,_=prepare_units([row])
    assert '不得' in prompt[0]['evidence_format'] and '距离对应' in prompt[0]['evidence_format']
    assert originals['E1'][0]['quote']=='0.4米'


def test_final_answer_preserves_original_ocr_label_and_source_url(monkeypatch):
    from datetime import datetime,timezone
    from app.campus_qa import campus_answer
    with session_factory()() as session:
        doc,chunks=add(session,'拍摄点布置示意图',['背景应为白色、蓝色，不得用绿色、红色。'])
        chunks[0].location='Word 图像 1 / OCR 1 conf=0.980'
        doc.source_url='https://jwc.njxzc.edu.cn/image.docx'
        session.commit()
        answer=campus_answer(session,'拍摄点布置示意图 背景颜色',datetime.now(timezone.utc),use_model=False)
        assert answer['sources'] and '【图像OCR候选，须核对原图】' in answer['answer']
        assert any(r['url']==doc.source_url and r['quote']==chunks[0].text for r in answer['sources'])
        assert any('箭头' in w for w in answer['warnings'])


@pytest.mark.parametrize('boundary',['private_parent','withdrawn_parent','published','existing_chunks','changed_sha'])
def test_image_native_repair_retains_review_boundaries(boundary,monkeypatch):
    import app.ingest as ingest
    def forbidden(_):raise AssertionError('Denied source must never reach OCR')
    monkeypatch.setattr(ingest,'local_ocr',forbidden)
    content=payload(document())
    with session_factory()() as session:
        parent,_=add(session,'公开通知',['附件'])
        doc,_=add(session,'图像Word',[],parent_id=parent.id)
        doc.mime='.docx';doc.status='parse_failed';doc.sha256=sha256(content).hexdigest()
        if boundary=='private_parent':parent.visibility='private'
        if boundary=='withdrawn_parent':parent.status='withdrawn'
        if boundary=='published':doc.status='published'
        if boundary=='existing_chunks':session.add(Chunk(id='existing',document_id=doc.id,text='原有审核内容',location='原位置'))
        if boundary=='changed_sha':doc.sha256='0'*64
        session.commit()
        with pytest.raises(ValueError):module.BusinessStore(session).parse_native_original(doc,content,'.docx')
        session.rollback()
        assert not list(session.scalars(select(Audit)))
