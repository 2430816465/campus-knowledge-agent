from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from zipfile import ZipFile
import pytest
from sqlalchemy import select
from app.file_formats import office_package_suffix
from app.ingest import parse
from app.db import Chunk,Audit
from scripts.convert_archive_doc import BusinessStore
from tests.test_appeal_retrieval import session_factory,add


def workbook_bytes():
    from openpyxl import Workbook
    workbook=Workbook();sheet=workbook.active;sheet.title='示范课安排'
    sheet.append(['课程名称','评议时间']);sheet.append(['教学方法','2026-10-20'])
    output=BytesIO();workbook.save(output);workbook.close();return output.getvalue()


def repack(payload, extra=None, modify=None):
    output=BytesIO()
    with ZipFile(BytesIO(payload)) as source,ZipFile(output,'w') as dest:
        for name in source.namelist():
            value=source.read(name)
            if modify:value=modify(name,value)
            dest.writestr(name,value)
        for name,value in (extra or {}).items():dest.writestr(name,value)
    return output.getvalue()


def test_mislabeled_excel_routes_by_package_and_preserves_source_declaration():
    payload=workbook_bytes()
    assert office_package_suffix(payload,'.doc')=='.xlsx'
    title,date,publisher,blocks,warnings,links=parse(payload,'.doc','https://jwc.njxzc.edu.cn/a.doc')
    assert blocks[1]==('工作表 示范课安排 行 2','教学方法 | 2026-10-20')
    assert date is None and title is None
    assert any('.doc' in w and '.xlsx' in w for w in warnings)
    assert office_package_suffix(b'legacy binary','.doc') is None
    assert office_package_suffix(payload,'.txt') is None


def test_real_word_package_detects_correctly():
    from docx import Document
    doc=Document();doc.add_paragraph('公开通知');stream=BytesIO();doc.save(stream)
    payload=stream.getvalue()
    assert office_package_suffix(payload,'.xls')=='.docx'
    assert parse(payload,'.xls')[3][0]==('Word 段落 1','公开通知')


@pytest.mark.parametrize('extra',[
    {'word/document.xml':b'<document/>'},
    {'../outside':b'unsafe'},
    {'xl/vbaProject.bin':b'macro'},
    {'ppt/presentation.xml':b'<presentation/>'},
])
def test_ambiguous_macro_and_unsafe_packages_fail_before_parser(extra):
    with pytest.raises(ValueError):office_package_suffix(repack(workbook_bytes(),extra),'.doc')


def test_invalid_declared_type_xml_and_budget_fail_closed():
    payload=workbook_bytes()
    def macro_type(name,value):
        return value.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                             b'application/vnd.ms-excel.sheet.macroEnabled.main+xml') if name=='[Content_Types].xml' else value
    with pytest.raises(ValueError,match='macro'):office_package_suffix(repack(payload,modify=macro_type),'.doc')
    with pytest.raises(ValueError,match='budget'):office_package_suffix(b'PK'+b'x'*(20*1024*1024),'.doc')
    with pytest.raises(ValueError,match='Malformed'):office_package_suffix(b'PKbroken','.doc')
    with pytest.raises(ValueError):office_package_suffix(repack(payload,modify=lambda n,v: b'<!DOCTYPE a><Types/>' if n=='[Content_Types].xml' else v),'.doc')


def source(session,payload):
    parent,_=add(session,'示范课观摩通知',['查看附件'])
    doc,chunks=add(session,'示范课安排.doc',[],parent_id=parent.id)
    doc.sha256=sha256(payload).hexdigest();doc.raw_name=doc.sha256+'.doc';doc.mime='.doc'
    doc.status='parse_failed';doc.error='old unsupported';doc.published_at='2026-10-01';session.commit()
    return parent,doc


def test_repair_same_original_without_word_or_embeddings_and_reuse_is_read_only():
    payload=workbook_bytes()
    with session_factory()() as session:
        parent,doc=source(session,payload);store=BusinessStore(session)
        before={k:getattr(doc,k) for k in ('id','sha256','raw_name','mime','source_url','parent_id','title','published_at','version')}
        store.parse_native_original(doc,payload,'.xlsx')
        assert {k:getattr(doc,k) for k in before}==before
        assert doc.status=='pending_review' and doc.error is None
        rows=list(session.scalars(select(Chunk).where(Chunk.document_id==doc.id)))
        assert len(rows)==2 and all(r.vector is None for r in rows)
        assert rows[1].text=='教学方法 | 2026-10-20'
        audits=list(session.scalars(select(Audit).where(Audit.action=='archive_native_format_parse')))
        assert len(audits)==1 and audits[0].changes['word_invoked'] is False
        assert store.parse_native_original(doc,payload,'.xlsx').id==doc.id
        assert len(list(session.scalars(select(Chunk).where(Chunk.document_id==doc.id))))==2
        assert len(list(session.scalars(select(Audit).where(Audit.action=='archive_native_format_parse'))))==1


@pytest.mark.parametrize('boundary',['private_parent','withdrawn_parent','published_source','existing_chunks','different_sha'])
def test_repair_denies_review_permission_existing_data_and_changed_bytes(boundary):
    payload=workbook_bytes()
    with session_factory()() as session:
        parent,doc=source(session,payload)
        if boundary=='private_parent':parent.visibility='private'
        if boundary=='withdrawn_parent':parent.status='withdrawn'
        if boundary=='published_source':doc.status='published'
        if boundary=='existing_chunks':session.add(Chunk(id='existing',document_id=doc.id,text='已有原文',location='原位置'))
        if boundary=='different_sha':doc.sha256='0'*64
        session.commit()
        with pytest.raises(ValueError):BusinessStore(session).parse_native_original(doc,payload,'.xlsx')
        session.rollback()
        assert not list(session.scalars(select(Audit).where(Audit.action=='archive_native_format_parse')))


def test_archive_native_route_never_invokes_word_and_persists_success(tmp_path):
    from tests.test_convert_archive_doc import module,archived,FakeStore
    class NativeStore(FakeStore):
        def parse_native_original(self,original,content,suffix):
            assert suffix=='.xlsx'
            original.status='pending_review';return original
    def forbidden(*args):raise AssertionError('Word must not run')
    row=archived(tmp_path,workbook_bytes())
    event=module.convert_row(row,NativeStore(),tmp_path/'school-archive',tmp_path/'out',tmp_path/'helper',converter=forbidden)
    assert event['state']=='native_parsed' and event['word_invoked'] is False
    assert event['converted_document_id']==event['original_document_id']
    # Main frontier success logic shares the pre-existing imported state contract.
    import sqlite3
    with sqlite3.connect(':memory:') as db:
        db.executescript('CREATE TABLE urls(url TEXT PRIMARY KEY,import_state TEXT,import_error TEXT,document_id TEXT);'
            'CREATE TABLE doc_conversions(url TEXT PRIMARY KEY,original_sha256 TEXT,state TEXT,attempts INTEGER,original_document_id TEXT,converted_document_id TEXT,converted_sha256 TEXT,checked_at TEXT,error TEXT);')
        db.execute('INSERT INTO urls(url) VALUES (?)',(row['url'],))
        module.save_frontier_event(db,event)
        assert db.execute('SELECT import_state,document_id FROM urls').fetchone()==('imported',event['original_document_id'])
